"""Export rendering: one dataset, three formats with identical records and source references."""

import io
import json
from typing import Any

import openpyxl
import pdfplumber

from attendance_ai.exports.datasets import RECORD_FIELDS, ExportDataset
from attendance_ai.exports.render import MEDIA_TYPES, excel_value, pdf_text, render


def _record(index: int, **overrides: Any) -> dict[str, Any]:
    record: dict[str, Any] = {
        "attendance_date": f"2026-08-{3 + index % 20:02d}",
        "employee_id": f"E0{10 + index % 9}",
        "employee_name": "Rahul Sharma",
        "department_id": "ENG",
        "department_name": "Engineering",
        "status": "PRESENT",
        "check_in": "09:15",
        "check_out": "18:24",
        "total_hours": 9.15,
        "remarks": None,
        "review_status": "auto_accepted",
        "extraction_method": "native_csv",
        "extraction_confidence": 1.0,
        "source_file": "acme_engineering_biometric_2026-08.csv",
        "source_page_or_row": f"row {index + 2}",
        "record_key": f"{index:016x}",
        "record_id": f"00000000-0000-0000-0000-{index:012d}",
    }
    record.update(overrides)
    return record


def _dataset(records: list[dict[str, Any]]) -> ExportDataset:
    meta = {
        "title": "Attendance records",
        "generated_at": "2026-09-29T10:15:00+00:00",
        "generated_by": "acme.admin",
        "role": "admin",
        "tenant_id": "acme",
        "tenant_name": "Acme Corp",
        "scope": "all of Acme Corp",
        "classification": "restricted",
        "data_version": 3,
        "knowledge_version": 1,
        "record_count": len(records),
        "filters": {},
    }
    return ExportDataset("records", "Attendance records", meta, records)


def test_every_format_carries_the_same_records_and_source_references() -> None:
    records = [_record(i) for i in range(40)]
    dataset = _dataset(records)
    expected = {(r["record_key"], r["source_file"], r["source_page_or_row"]) for r in records}

    as_json = json.loads(render(dataset, "json").content)
    assert {
        (r["record_key"], r["source_file"], r["source_page_or_row"]) for r in as_json["records"]
    } == expected
    assert as_json["export"]["record_count"] == 40

    sheet = openpyxl.load_workbook(io.BytesIO(render(dataset, "xlsx").content))["Records"]
    header = [cell.value for cell in sheet[1]]
    assert header == [label for _, label in RECORD_FIELDS]
    key, source, location = (header.index(name) for name in ("Record key", "Source file", "Source location"))
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert {(row[key], row[source], row[location]) for row in rows} == expected

    with pdfplumber.open(io.BytesIO(render(dataset, "pdf").content)) as pdf:
        text = "".join(page.extract_text() or "" for page in pdf.pages)
    assert all(record["record_key"] in text for record in records)
    assert all(record["source_page_or_row"] in text for record in records)


def test_media_types_and_file_names() -> None:
    dataset = _dataset([_record(0)])
    for fmt in ("json", "xlsx", "pdf"):
        rendered = render(dataset, fmt)  # type: ignore[arg-type]
        assert rendered.media_type == MEDIA_TYPES[fmt]
        assert rendered.filename == f"attendance-records_acme_20260929-101500.{fmt}"
    assert render(dataset, "pdf").content.startswith(b"%PDF-")


def test_spreadsheet_cells_are_never_formulas() -> None:
    for text in ('=HYPERLINK("http://evil","x")', "+1", "-2+3", "@SUM(A1)"):
        assert excel_value(text) == "'" + text
    assert excel_value("Present") == "Present"
    assert excel_value(9.15) == 9.15
    dataset = _dataset([_record(0, remarks='=HYPERLINK("http://evil","x")')])
    sheet = openpyxl.load_workbook(io.BytesIO(render(dataset, "xlsx").content))["Records"]
    remarks = [cell.value for cell in sheet[1]].index("Remarks") + 1
    cell = sheet.cell(row=2, column=remarks)
    assert cell.data_type == "s" and str(cell.value).startswith("'=")


def test_pdf_text_is_escaped_so_uploaded_content_cannot_inject_markup() -> None:
    assert pdf_text("<b>bold</b> & more\nnext") == "&lt;b&gt;bold&lt;/b&gt; &amp; more<br/>next"
    dataset = _dataset([_record(0, remarks="<font size=40>big</font> & <b>")])
    assert render(dataset, "pdf").content.startswith(b"%PDF-")  # renders instead of failing on markup


def test_query_exports_include_the_question_answer_and_citations() -> None:
    dataset = ExportDataset(
        "query",
        "Question and answer",
        {**_dataset([]).meta, "request_id": "req_0123456789abcdef"},
        [_record(0)],
        query={
            "request_id": "req_0123456789abcdef",
            "question": "What was attendance?",
            "answer": "Attendance was 92.41% [D1].",
            "outcome": "answered",
            "confidence": {"score": 0.9, "band": "high", "reasons": []},
            "citations": [{"id": "D1", "kind": "records", "source_file": "a.csv", "locations": "rows 2-4"}],
        },
        result=[{"attendance_pct": 92.41, "counted_days": 79}],
    )
    body = json.loads(render(dataset, "json").content)
    assert body["query"]["answer"] == "Attendance was 92.41% [D1]."
    assert body["result"] == [{"attendance_pct": 92.41, "counted_days": 79}]
    workbook = openpyxl.load_workbook(io.BytesIO(render(dataset, "xlsx").content))
    assert {"About", "Answer", "Citations", "Result", "Records"} <= set(workbook.sheetnames)
    with pdfplumber.open(io.BytesIO(render(dataset, "pdf").content)) as pdf:
        text = "".join(page.extract_text() or "" for page in pdf.pages)
    assert "92.41" in text and "a.csv" in text and record_key_of(0) in text


def record_key_of(index: int) -> str:
    return f"{index:016x}"
