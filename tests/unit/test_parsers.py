"""File parsers, checked row for row against the sample files' ground truth, plus small generated files
for the edge cases.

Parsers only extract: the raw text of each canonical field, and the exact place each row came from.
Dates and times are compared after the normaliser's own parsing, because every format writes them
differently.
"""

import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pytesseract
import pytest
from PIL import Image

from attendance_ai.ingestion.normalize import KEY_FIELDS, parse_date, parse_time
from attendance_ai.ingestion.parsers import (
    CorruptFileError,
    OcrUnavailableError,
    UnsupportedFileError,
    parse_file,
)
from attendance_ai.ingestion.types import (
    MEDIA_CSV,
    MEDIA_DOCX,
    MEDIA_PDF,
    MEDIA_XLSX,
    ExtractedRow,
    ParseResult,
)

from ..support import PROJECT_ROOT, find_tesseract

SAMPLES = PROJECT_ROOT / "sample_data" / "tenants"
SCANNED = "acme_hr_register_scanned_2026-08.pdf"
TESSERACT = find_tesseract()
REVIEW_THRESHOLD = 0.8  # the seed's ocr_review_threshold: below it, a key field sends the record to review
MEDIA = {".csv": MEDIA_CSV, ".xlsx": MEDIA_XLSX, ".docx": MEDIA_DOCX, ".pdf": MEDIA_PDF}

needs_tesseract = pytest.mark.skipif(TESSERACT is None, reason="Tesseract is not installed")


def _parse(path: Path, **options: Any) -> ParseResult:
    return parse_file(path, MEDIA[path.suffix], **options)


def _ground_truth(tenant: str, filename: str) -> dict[str, dict[str, Any]]:
    path = SAMPLES / tenant / "ground_truth" / "canonical_records.jsonl"
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    return {row["source_page_or_row"]: row for row in rows if row["source_file"] == filename}


def _by_locator(result: ParseResult) -> dict[str, ExtractedRow]:
    return {row.locator.describe(): row for row in result.rows}


def _clock(raw: str | None) -> str | None:
    parsed = parse_time(raw)
    return parsed.strftime("%H:%M") if parsed else None


def _mismatches(row: ExtractedRow, truth: dict[str, Any]) -> list[str]:
    values = row.values
    checks = {
        "employee_id": (values.get("employee_id"), truth["employee_id"]),
        "date": (parse_date(values.get("date")), date.fromisoformat(truth["attendance_date"])),
        "status": (values.get("status"), truth["raw_status"]),
        "check_in": (_clock(values.get("check_in")), truth["check_in"]),
        "check_out": (_clock(values.get("check_out")), truth["check_out"]),
    }
    return [f"{name}: {got!r} != {want!r}" for name, (got, want) in checks.items() if got != want]


def _blank_scan(path: Path) -> Path:
    """A one-page PDF holding only an image of blank paper: no text layer, so it goes to OCR."""
    Image.new("L", (1275, 1650), 255).save(path, "PDF", resolution=150)
    return path


# --- the sample files ---------------------------------------------------------------------------------

NATIVE = [
    ("acme", "acme_engineering_biometric_2026-08.csv"),
    ("acme", "acme_sales_muster_2026-08.xlsx"),
    ("acme", "acme_operations_weekly_report_2026-08.docx"),
    ("acme", "acme_finance_attendance_2026-08.pdf"),
    ("globex", "globex_engineering_biometric_2026-08.csv"),
    ("globex", "globex_support_muster_2026-08.xlsx"),
]


@pytest.mark.parametrize(("tenant", "filename"), NATIVE)
def test_native_files_match_the_ground_truth_row_for_row(tenant: str, filename: str) -> None:
    expected = _ground_truth(tenant, filename)
    result = _parse(SAMPLES / tenant / "inputs" / filename)
    found = _by_locator(result)
    assert len(found) == len(result.rows), "two rows claim the same place in the file"
    assert set(found) == set(expected)
    for locator, truth in expected.items():
        row = found[locator]
        assert row.method == truth["extraction_method"], locator
        assert row.confidence == {}, locator  # native extraction is exact
        assert _mismatches(row, truth) == [], locator
        assert row.values.get("remarks") == truth["remarks"], locator


def test_a_memo_without_a_table_gives_narrative_only() -> None:
    memo = SAMPLES / "acme" / "scenarios" / "prompt_injection" / "acme_operations_memo_2026-08.docx"
    result = _parse(memo)
    assert result.rows == []
    assert any("no attendance table" in warning for warning in result.warnings)
    text = " ".join(block.text for block in result.texts)
    assert "offsite" in text and "safety training" in text
    # The embedded instructions are extracted as plain data; the chunker quarantines them later.
    assert "Ignore all previous instructions" in text
    assert all(block.locator.type == "paragraph" for block in result.texts)


@pytest.fixture(scope="module")
def scanned() -> ParseResult:
    if TESSERACT is None:
        pytest.skip("Tesseract is not installed")
    return _parse(SAMPLES / "acme" / "inputs" / SCANNED, tesseract_cmd=TESSERACT)


def test_scanned_register_rows_are_read_by_ocr_in_their_places(scanned: ParseResult) -> None:
    expected = _ground_truth("acme", SCANNED)
    found = _by_locator(scanned)
    assert {row.method for row in scanned.rows} == {"ocr"}
    placed = [
        locator
        for locator, truth in expected.items()
        if locator in found and found[locator].values.get("employee_id") == truth["employee_id"]
    ]
    assert len(placed) >= 0.95 * len(expected), sorted(set(expected) - set(placed))


def test_scanned_register_flags_the_degraded_cells_for_review(scanned: ParseResult) -> None:
    found = _by_locator(scanned)
    for locator in ("page 2, row 7", "page 3, row 10"):  # a smudged status cell and a faint check-in
        row = found.get(locator)
        if row is None:
            continue  # a row that could not be read at all is not a guess either
        lowest = min(row.confidence.get(name, 1.0) for name in KEY_FIELDS)
        assert lowest < REVIEW_THRESHOLD, (locator, row.values, row.confidence)


def test_confident_ocr_values_are_correct(scanned: ParseResult) -> None:
    """Anything the OCR is sure about is accepted without review, so it must be right."""
    expected = _ground_truth("acme", SCANNED)
    for locator, row in _by_locator(scanned).items():
        truth = expected.get(locator)
        confident = all(row.confidence.get(name, 1.0) >= REVIEW_THRESHOLD for name in KEY_FIELDS)
        if truth is not None and confident:
            assert _mismatches(row, truth) == [], (locator, row.values, row.confidence)


# --- CSV ----------------------------------------------------------------------------------------------


def test_csv_rows_keep_their_physical_line_numbers(tmp_path: Path) -> None:
    path = tmp_path / "export.csv"
    path.write_bytes(
        b"Attendance export - August 2026,,,,\n"
        b"Emp ID,Name,Date,Status,Check In\n"
        b"E001,Rahul Sharma,03/08/2026,Present,09:15\n"
        b",,,,\n"
        b'E002,"Ananya\nIyer",03/08/2026,Absent,\n'  # a quoted value spanning two lines
        b"Emp ID,Name,Date,Status,Check In\n"  # the header repeated inside the data
        b"E003,Vikram Reddy,04/08/2026,Leave,\n"
        b",Total present,,1,\n"  # a footer: no employee ID and no date
    )
    result = _parse(path)
    assert [row.locator.describe() for row in result.rows] == ["row 3", "row 5", "row 8"]
    assert [row.values["employee_id"] for row in result.rows] == ["E001", "E002", "E003"]
    assert result.rows[0].values["check_in"] == "09:15"
    assert result.rows[1].values["check_in"] is None
    assert {row.method for row in result.rows} == {"native_csv"}


def test_csv_with_a_byte_order_mark_and_semicolons(tmp_path: Path) -> None:
    path = tmp_path / "export.csv"
    path.write_bytes("﻿Emp ID;Name;Date;Status\r\nE001;Rahul Sharma;2026-08-03;Present\r\n".encode())
    result = _parse(path)
    assert [(r.locator.describe(), r.values["employee_id"], r.values["status"]) for r in result.rows] == [
        ("row 2", "E001", "Present")
    ]


def test_csv_without_a_recognisable_header_gives_a_warning_and_no_rows(tmp_path: Path) -> None:
    path = tmp_path / "export.csv"
    path.write_bytes(b"a,b,c\n1,2,3\n")
    result = _parse(path)
    assert result.rows == []
    assert result.warnings


def test_binary_data_is_not_read_as_csv(tmp_path: Path) -> None:
    path = tmp_path / "export.csv"
    path.write_bytes(b"Emp ID,Name\x00\x01\x02")
    with pytest.raises(CorruptFileError):
        _parse(path)


def test_unknown_media_types_are_refused(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_bytes(b"hello")
    with pytest.raises(UnsupportedFileError):
        parse_file(path, "text/plain")


# --- XLSX ---------------------------------------------------------------------------------------------


def test_xlsx_long_layout_gives_one_row_per_line_with_its_cell(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Log"
    sheet.append(["Attendance log - August 2026"])
    sheet.append([])
    sheet.append(["Emp ID", "Name", "Date", "Status", "In Time", "Out Time"])
    sheet.append(["E001", "Rahul Sharma", datetime(2026, 8, 3), "Present", time(9, 15), time(18, 5)])
    sheet.append(["E002", "Ananya Iyer", datetime(2026, 8, 3), "Absent", None, None])
    path = tmp_path / "log.xlsx"
    workbook.save(path)

    result = _parse(path)
    assert [row.locator.describe() for row in result.rows] == ["sheet 'Log', cell A4", "sheet 'Log', cell A5"]
    first = result.rows[0].values
    assert (first["employee_id"], first["date"], first["status"]) == ("E001", "2026-08-03", "Present")
    assert (first["check_in"], first["check_out"]) == ("09:15", "18:05")
    assert result.rows[1].values["check_in"] is None


def test_xlsx_muster_matrix_takes_the_month_from_the_title_above_it(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Muster"
    sheet.append(["Sales muster - August 2026"])
    sheet.append([])
    sheet.append(["Emp ID", "Employee Name", 1, 2, 3, 4, 5, 6])
    sheet.append(["E005", "Arjun Mehta", "WO", "WO", "P", "A", None, "L"])
    sheet.append([])
    sheet.append(["Legend: P = Present, A = Absent"])
    workbook.create_sheet("Legend").append(["P", "Present"])
    path = tmp_path / "muster.xlsx"
    workbook.save(path)

    result = _parse(path)
    assert [(r.locator.describe(), r.values["date"], r.values["status"]) for r in result.rows] == [
        ("sheet 'Muster', cell C4", "2026-08-01", "WO"),
        ("sheet 'Muster', cell D4", "2026-08-02", "WO"),
        ("sheet 'Muster', cell E4", "2026-08-03", "P"),
        ("sheet 'Muster', cell F4", "2026-08-04", "A"),
        ("sheet 'Muster', cell H4", "2026-08-06", "L"),  # the empty day 5 is not a record
    ]
    assert {row.values["employee_id"] for row in result.rows} == {"E005"}
    assert any("'Legend'" in warning for warning in result.warnings)  # a sheet without a table is skipped


def test_a_damaged_workbook_is_reported_as_corrupt(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"PK\x03\x04 this is not a real workbook")
    with pytest.raises(CorruptFileError):
        _parse(path)


# --- DOCX ---------------------------------------------------------------------------------------------


def test_docx_tables_become_rows_and_paragraphs_narrative_under_their_heading(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("Weekly report", level=0)
    document.add_heading("Summary", level=1)
    document.add_paragraph("Attendance was steady this week.")
    table = document.add_table(rows=3, cols=4)
    lines = [
        ("Emp ID", "Date", "Status", "Remarks"),
        ("E009", "03/08/2026", "Half Day", "Half day - bank work"),
        ("E010", "03/08/2026", "Present", ""),
    ]
    for row, values in zip(table.rows, lines, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_heading("Notes", level=1)
    document.add_paragraph("Two employees were on leave.")
    path = tmp_path / "report.docx"
    document.save(str(path))

    result = _parse(path)
    assert [(r.locator.describe(), r.values["employee_id"], r.values["remarks"]) for r in result.rows] == [
        ("table 1, row 1", "E009", "Half day - bank work"),
        ("table 1, row 2", "E010", None),
    ]
    # Paragraphs are numbered over the non-empty body paragraphs, headings included.
    assert [(t.text, t.locator.describe(), t.heading) for t in result.texts] == [
        ("Attendance was steady this week.", "paragraph 3", "Summary"),
        ("Two employees were on leave.", "paragraph 5", "Notes"),
    ]


def test_a_damaged_word_document_is_reported_as_corrupt(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a word document")
    with pytest.raises(CorruptFileError):
        _parse(path)


# --- PDF ----------------------------------------------------------------------------------------------


def test_a_damaged_pdf_is_reported_as_corrupt(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.7\nthis is not a real pdf\n")
    with pytest.raises(CorruptFileError):
        _parse(path)


@needs_tesseract
def test_a_blank_scanned_page_is_reported_not_guessed(tmp_path: Path) -> None:
    result = _parse(_blank_scan(tmp_path / "blank.pdf"), tesseract_cmd=TESSERACT)
    assert result.rows == []
    assert "page 1: blank page" in result.warnings


def test_a_scan_without_tesseract_fails_with_a_clear_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # The parser sets pytesseract's module-level command; restore it after this test.
    monkeypatch.setattr(pytesseract.pytesseract, "tesseract_cmd", pytesseract.pytesseract.tesseract_cmd)
    missing = tmp_path / "missing" / "tesseract.exe"
    with pytest.raises(OcrUnavailableError):
        _parse(_blank_scan(tmp_path / "scan.pdf"), tesseract_cmd=str(missing))
