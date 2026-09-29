"""XLSX (openpyxl, cached formula values). Two layouts are recognised on each sheet:

- long: one row per employee per day under a header row naming at least 3 known fields;
- matrix (muster): an Emp ID and/or Name column plus one column per day, whose header cells are real
  dates, or day numbers 1-31 with the month taken from a title above the table (for example
  "August 2026") or from the sheet name. Each non-empty day cell is one row.

Sheets without either (a legend, a summary) are skipped.
"""

from __future__ import annotations

import calendar
import re
from collections.abc import Sequence
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from attendance_ai.ingestion.parsers import (
    HEADER_SEARCH_ROWS,
    MIN_HEADER_FIELDS,
    CorruptFileError,
    clean,
    header_columns,
    is_record,
    repeats_header,
)
from attendance_ai.ingestion.types import ExtractedRow, Locator, ParseResult

MIN_MATRIX_DAYS = 5
IDENTITY_FIELDS = ("employee_id", "employee_name", "department")
TIME_FIELDS = ("check_in", "check_out")

_MONTHS = {name.lower(): number for number, name in enumerate(calendar.month_abbr) if name}
_MONTH_YEAR = re.compile(
    r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?[\s,/'-]*(\d{4})\b", re.I
)
_NUMERIC_MONTH_YEAR = re.compile(r"\b(\d{1,2})[/.-](\d{4})\b|\b(\d{4})-(\d{1,2})\b")

_Rows = list[list[object]]


def parse_xlsx(path: Path) -> ParseResult:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CorruptFileError(f"The file is not a readable Excel workbook: {exc}") from exc
    result = ParseResult()
    try:
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()  # read-only sheets trust the stored size, which some writers get wrong
            rows = [list(row) for row in sheet.iter_rows(min_row=1, min_col=1, values_only=True)]
            _parse_sheet(sheet.title, rows, result)
    except (KeyError, ValueError, SyntaxError, OSError) as exc:
        raise CorruptFileError(f"The Excel workbook is damaged: {exc}") from exc
    finally:
        workbook.close()
    if not result.rows:
        result.warnings.append("no attendance table found in any sheet")
    return result


def _parse_sheet(title: str, rows: _Rows, result: ParseResult) -> None:
    for index, cells in enumerate(rows[:HEADER_SEARCH_ROWS]):
        columns = header_columns(cells)
        identity = {i: name for i, name in columns.items() if name in IDENTITY_FIELDS}
        days = _day_columns(cells, rows[:index], title, result) if identity else None
        if days:
            _matrix_rows(title, rows, index, identity, days, result)
            return
        if len(columns) >= MIN_HEADER_FIELDS:
            _long_rows(title, rows, index, columns, result)
            return
    if any(any(clean(value) for value in cells) for cells in rows):
        result.warnings.append(f"sheet '{title}': no attendance table found; skipped")


def _long_rows(
    title: str, rows: _Rows, header_index: int, columns: dict[int, str], result: ParseResult
) -> None:
    first_column = next(i for i, value in enumerate(rows[header_index]) if clean(value) is not None)
    for row_no, cells in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if repeats_header(cells, columns):
            continue
        values = {name: _cell_text(_at(cells, i), name) for i, name in columns.items()}
        if is_record(values):
            cell = f"{get_column_letter(first_column + 1)}{row_no}"
            locator = Locator("cell", sheet=title, cell=cell, row=row_no)
            result.rows.append(ExtractedRow(values, locator, "native_xlsx"))


def _matrix_rows(
    title: str,
    rows: _Rows,
    header_index: int,
    identity: dict[int, str],
    days: dict[int, date],
    result: ParseResult,
) -> None:
    for row_no, cells in enumerate(rows[header_index + 1 :], start=header_index + 2):
        who = {name: _cell_text(_at(cells, i), name) for i, name in identity.items()}
        if not (who.get("employee_id") or who.get("employee_name")):
            continue
        for column, day in days.items():
            code = _cell_text(_at(cells, column), "status")
            if code is None:
                continue
            values = {**who, "date": day.isoformat(), "status": code}
            cell = f"{get_column_letter(column + 1)}{row_no}"
            result.rows.append(
                ExtractedRow(values, Locator("cell", sheet=title, cell=cell, row=row_no), "native_xlsx")
            )


def _day_columns(
    cells: Sequence[object], above: _Rows, title: str, result: ParseResult
) -> dict[int, date] | None:
    """Column index -> date for a matrix header row, or None if the row is not one."""
    dates = {i: _as_date(value) for i, value in enumerate(cells) if isinstance(value, date)}
    if len(dates) >= MIN_MATRIX_DAYS:
        return dates
    numbers = {i: n for i, value in enumerate(cells) if (n := _day_number(value)) is not None}
    if len(numbers) < MIN_MATRIX_DAYS:
        return None
    period = _period(above, title)
    if period is None:
        result.warnings.append(f"sheet '{title}': day-number columns but no month and year above the table")
        return None
    year, month = period
    last_day = calendar.monthrange(year, month)[1]
    return {i: date(year, month, n) for i, n in numbers.items() if n <= last_day}


def _period(above: _Rows, title: str) -> tuple[int, int] | None:
    """(year, month) from a title above the table ("August 2026", "08/2026"), else the sheet name, else a
    date cell above the table."""
    texts = [value for cells in above for value in cells if isinstance(value, str)]
    for text in [*texts, title]:
        if found := _month_in(text):
            return found
    dates = [value for cells in above for value in cells if isinstance(value, date)]
    return (dates[0].year, dates[0].month) if dates else None


def _month_in(text: str) -> tuple[int, int] | None:
    if match := _MONTH_YEAR.search(text):
        year, month = int(match[2]), _MONTHS[match[1].lower()]
    elif match := _NUMERIC_MONTH_YEAR.search(text):
        year, month = (int(match[2]), int(match[1])) if match[1] else (int(match[3]), int(match[4]))
    else:
        return None
    return (year, month) if 1900 <= year <= 2999 and 1 <= month <= 12 else None


def _day_number(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, str) and value.strip().isdigit():
        value = int(value.strip())
    return value if isinstance(value, int) and 1 <= value <= 31 else None


def _as_date(value: date) -> date:
    return value.date() if isinstance(value, datetime) else value


def _at(cells: Sequence[object], index: int) -> object:
    return cells[index] if index < len(cells) else None


def _cell_text(value: object, name: str) -> str | None:
    """A cell as raw text: dates as ISO YYYY-MM-DD, times as HH:MM, whole numbers without ".0"."""
    if isinstance(value, datetime):
        if name in TIME_FIELDS:
            return _clock(value.time())
        if name == "date" or value.time() == time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return _clock(value)
    if isinstance(value, timedelta):
        minutes = round(value.total_seconds() / 60)
        return f"{minutes // 60}:{minutes % 60:02d}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def _clock(value: time) -> str:
    return f"{value:%H:%M:%S}" if value.second else f"{value:%H:%M}"
